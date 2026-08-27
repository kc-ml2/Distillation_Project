import unittest
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from PIL import Image

from docs.experiments.plot_epoch_metrics import CAPTION, MODELS, RETRIEVAL, at_epoch_ends, plot_excel, write_excel


class EpochEndSeriesTest(unittest.TestCase):
    def test_selects_epoch_end_steps_in_checkpoint_order(self):
        points = {500: 0.1, 37346: 1.0, 74692: 2.0}

        self.assertEqual(at_epoch_ends(points, epochs=2), [1.0, 2.0])

    def test_rejects_missing_epoch(self):
        with self.assertRaisesRegex(ValueError, "74692"):
            at_epoch_ends({37346: 1.0}, epochs=2)

    def test_writes_epoch_19_performance_relative_to_base(self):
        models = {
            "Base Capacity": 200.0,
            "Small Baseline": 100.0,
            "LM Distillation": 150.0,
            "Distill All": 180.0,
        }
        data = {}
        for sheet, metrics in (("ITC Retrieval", RETRIEVAL),
                               ("ITM Retrieval", RETRIEVAL),
                               ("Caption", CAPTION)):
            data[sheet] = {
                model: {metric: [value] * 20 for metric in metrics}
                for model, value in models.items()
            }
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "metrics.xlsx"
            write_excel(data, path)
            rows = list(load_workbook(path, data_only=True)["Final Base Relative (%)"].values)

        cider = next(row for row in rows[1:] if row[:2] == ("Caption", "CIDEr"))
        self.assertEqual(cider, ("Caption", "CIDEr", 200.0, 100.0, 50.0, 180.0, 90.0))

    def test_exports_presentation_pngs(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook = Workbook()
            workbook.remove(workbook.active)
            sheets = {
                "ITC Retrieval": ["txt_r1", "txt_r5", "txt_r10", "txt_r_mean", "img_r1", "img_r5", "img_r10", "img_r_mean", "r_mean"],
                "ITM Retrieval": ["txt_r1", "txt_r5", "txt_r10", "txt_r_mean", "img_r1", "img_r5", "img_r10", "img_r_mean", "r_mean"],
                "Caption": ["Bleu_1", "Bleu_2", "Bleu_3", "Bleu_4", "METEOR", "ROUGE_L", "CIDEr", "SPICE"],
            }
            for name, metrics in sheets.items():
                sheet = workbook.create_sheet(name)
                sheet.append(["Model", "Checkpoint Epoch", *metrics])
                for model in ("Base Capacity", "Small Baseline", "LM Distillation", "Distill All"):
                    sheet.append([model, 0, *([1.0] * len(metrics))])
                    sheet.append([model, 1, *([2.0] * len(metrics))])
            source = root / "metrics.xlsx"
            workbook.save(source)

            plot_excel(source, root / "overview.png")

            plots = list((root / "per_metric_plots").glob("*.png"))
            self.assertEqual(len(plots), 26)
            cider = root / "per_metric_plots/caption_cider.png"
            self.assertTrue(cider.exists())

            def dark_pixels(path):
                image = Image.open(path).convert("L")
                width, height = image.size
                crop = image.crop((2 * width // 3, 2 * height // 3, width, height))
                return sum(pixel < 64 for pixel in crop.get_flattened_data())

            self.assertGreater(dark_pixels(root / "overview.png"), 50_000)
            self.assertGreater(dark_pixels(cider), 16_000)

            comparison = root / "base_vs_distill_all"
            selected = {name: MODELS[name] for name in ("Base Capacity", "Distill All")}
            plot_excel(source, comparison / "all_metrics.png",
                       models=selected, per_metric=comparison)
            self.assertEqual(len(list(comparison.glob("*.png"))), 27)


if __name__ == "__main__":
    unittest.main()
