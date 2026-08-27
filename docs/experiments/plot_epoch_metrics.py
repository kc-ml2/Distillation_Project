"""Export every per-epoch metric to Excel, then plot that workbook."""

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from tensorboard.backend.event_processing.event_accumulator import EventAccumulator


ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
STEPS_PER_EPOCH = 37346
MODELS = {
    "Base Capacity": ("pt_checkpoint_base_logitscale_nodecay", "base_nodecay", "#4c78a8"),
    "Small Baseline": ("pt_smallreg_minilm_baseline", "small_solo", "#54a24b"),
    "LM Distillation": ("pt_smallreg_minilm_lm_distill", "small_distill", "#f58518"),
    "Distill All": ("pt_distill_all", None, "#e45756"),
}
RETRIEVAL = ["txt_r1", "txt_r5", "txt_r10", "txt_r_mean",
             "img_r1", "img_r5", "img_r10", "img_r_mean", "r_mean"]
CAPTION = ["Bleu_1", "Bleu_2", "Bleu_3", "Bleu_4", "METEOR", "ROUGE_L", "CIDEr", "SPICE"]


def at_epoch_ends(points, epochs=20):
    steps = [(epoch + 1) * STEPS_PER_EPOCH for epoch in range(epochs)]
    missing = [step for step in steps if step not in points]
    if missing:
        raise ValueError(f"missing epoch-end steps: {missing}")
    return [points[step] for step in steps]


def load_scalars(run):
    data = {}
    root = ROOT / "output" / run / "tensorboard"
    for event_file in sorted(root.rglob("events.out.tfevents*")):
        accumulator = EventAccumulator(str(event_file), size_guidance={"scalars": 0})
        accumulator.Reload()
        for tag in accumulator.Tags().get("scalars", []):
            data.setdefault(tag, {}).update({event.step: event.value for event in accumulator.Scalars(tag)})
    return data


def collect():
    data = {"ITC Retrieval": {}, "ITM Retrieval": {}, "Caption": {}}
    for model, (run, caption_prefix, _) in MODELS.items():
        scalars = load_scalars(run)
        for sheet, tag_group in (("ITC Retrieval", "val_retrieval_itc"),
                                 ("ITM Retrieval", "val_retrieval_itm")):
            data[sheet][model] = {
                metric: at_epoch_ends(scalars[f"{tag_group}/{metric}"])
                for metric in RETRIEVAL
            }
        if caption_prefix:
            data["Caption"][model] = {metric: [] for metric in CAPTION}
            for epoch in range(20):
                path = ROOT / f"output/caption_zeroshot/{caption_prefix}_ep{epoch:02d}/val_metrics_with_spice.json"
                scores = json.loads(path.read_text())
                for metric in CAPTION:
                    data["Caption"][model][metric].append(scores[metric])
        else:
            data["Caption"][model] = {
                metric: at_epoch_ends(scalars[f"val_caption/{metric}"])
                for metric in CAPTION
            }
    return data


def write_excel(data, path):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet, metrics in (("ITC Retrieval", RETRIEVAL), ("ITM Retrieval", RETRIEVAL), ("Caption", CAPTION)):
        worksheet = workbook.create_sheet(sheet)
        worksheet.append(["Model", "Checkpoint Epoch", *metrics])
        for model in MODELS:
            for epoch in range(20):
                worksheet.append([model, epoch, *(data[sheet][model][metric][epoch] for metric in metrics)])
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="3B5B92")
        worksheet.freeze_panes = "C2"
        worksheet.column_dimensions["A"].width = 20
        worksheet.column_dimensions["B"].width = 18
        for column in worksheet.iter_cols(min_col=3, max_col=2 + len(metrics)):
            worksheet.column_dimensions[column[0].column_letter].width = 14
            for cell in column[1:]:
                cell.number_format = "0.0000"
    relative = workbook.create_sheet("Final Base Relative (%)")
    relative.append(["Group", "Metric", "Base Capacity", "Small Baseline",
                     "Small vs Base (%)", "Distill All", "Distill All vs Base (%)"])
    for sheet, metrics in (("ITC Retrieval", RETRIEVAL), ("ITM Retrieval", RETRIEVAL),
                           ("Caption", CAPTION)):
        for metric in metrics:
            base = data[sheet]["Base Capacity"][metric][19]
            small = data[sheet]["Small Baseline"][metric][19]
            distilled = data[sheet]["Distill All"][metric][19]
            relative.append([sheet, metric, base, small, small / base * 100,
                             distilled, distilled / base * 100])
    for cell in relative[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3B5B92")
    relative.freeze_panes = "C2"
    relative.column_dimensions["A"].width = 18
    relative.column_dimensions["B"].width = 16
    for column in relative.iter_cols(min_col=3, max_col=7):
        relative.column_dimensions[column[0].column_letter].width = 22
        for cell in column[1:]:
            cell.number_format = "0.00"
    sources = workbook.create_sheet("Sources")
    sources.append(["Model", "TensorBoard run", "Caption source"])
    for model, (run, prefix, _) in MODELS.items():
        sources.append([model, f"output/{run}/tensorboard", "TensorBoard" if prefix is None else f"output/caption_zeroshot/{prefix}_epXX"])
    sources.append([])
    sources.append(["Checkpoint Epoch 0-19 maps to TensorBoard epoch-end steps 37346-746920."])
    workbook.save(path)


def plot_excel(path, output, models=MODELS, per_metric=None):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    workbook = load_workbook(path, data_only=True, read_only=True)
    groups = (("ITC Retrieval", RETRIEVAL), ("ITM Retrieval", RETRIEVAL), ("Caption", CAPTION))
    per_metric = output.parent / "per_metric_plots" if per_metric is None else per_metric
    per_metric.mkdir(parents=True, exist_ok=True)
    figure, axes = plt.subplots(9, 3, figsize=(20, 32), sharex=True)
    for column, (sheet, metrics) in enumerate(groups):
        rows = list(workbook[sheet].values)
        headers = list(rows[0])
        for row, metric in enumerate(metrics):
            axis = axes[row, column]
            metric_column = headers.index(metric)
            metric_figure, metric_axis = plt.subplots(figsize=(12, 7))
            for model, (_, _, color) in models.items():
                selected = [values for values in rows[1:] if values[0] == model]
                epochs = [values[1] for values in selected]
                scores = [values[metric_column] for values in selected]
                axis.plot(epochs, scores,
                          color=color, marker="o", markersize=2.5, linewidth=1.6, label=model)
                metric_axis.plot(epochs, scores, color=color, marker="o", markersize=6,
                                 linewidth=2.8, label=model)
            axis.set_title(metric)
            axis.grid(alpha=0.25)
            axis.set_xticks([0, 4, 8, 12, 16, 19])
            metric_axis.set_title(f"{sheet} — {metric}", fontsize=22, pad=16)
            metric_axis.set_xlabel("Checkpoint epoch", fontsize=17)
            metric_axis.set_ylabel(metric, fontsize=17)
            metric_axis.tick_params(labelsize=14)
            metric_axis.set_xticks([0, 4, 8, 12, 16, 19])
            metric_axis.grid(alpha=0.3)
            metric_axis.legend(fontsize=28, frameon=True, loc="lower right", markerscale=1.5,
                               borderpad=0.8, labelspacing=0.8)
            metric_figure.tight_layout()
            filename = f"{sheet.lower().replace(' ', '_')}_{metric.lower()}.png"
            metric_figure.savefig(per_metric / filename, dpi=200, bbox_inches="tight")
            plt.close(metric_figure)
        axes[0, column].set_title(f"{sheet}\n{metrics[0]}")
        for row in range(len(metrics), 9):
            axes[row, column].axis("off")
        axes[len(metrics) - 1, column].set_xlabel("Checkpoint epoch")
    figure.suptitle("Per-Epoch Validation Metrics", fontsize=20, y=0.998)
    handles, labels = axes[0, 0].get_legend_handles_labels()
    legend_axis = axes[8, 2]
    legend_axis.axis("off")
    legend_axis.legend(handles, labels, loc="center", fontsize=40, markerscale=2.5,
                       title="Models", title_fontsize=44, frameon=True,
                       borderpad=1.3, labelspacing=1.3)
    figure.tight_layout(rect=(0, 0, 1, 0.99))
    figure.savefig(output, dpi=180, bbox_inches="tight")
    plt.close(figure)


def main():
    excel = OUT / "per_epoch_all_metrics.xlsx"
    plot = OUT / "per_epoch_all_metrics.png"
    write_excel(collect(), excel)
    plot_excel(excel, plot)
    print(excel)
    print(plot)


if __name__ == "__main__":
    main()
