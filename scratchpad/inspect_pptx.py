from pptx import Presentation
p=Presentation("발표자료/파워포인트_템플릿.pptx")
print("slide size:", round(p.slide_width/914400,2),"x",round(p.slide_height/914400,2),"in")
print("existing slides:", len(p.slides))
print("\n=== layouts ===")
for i,layout in enumerate(p.slide_layouts):
    print(f"[{i}] {layout.name!r}")
    for ph in layout.placeholders:
        pf=ph.placeholder_format
        print(f"      idx={pf.idx} type={pf.type} name={ph.name!r}")
print("\n=== existing slides (layout + text preview) ===")
for i,s in enumerate(p.slides):
    txt=[]
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            txt.append(sh.text_frame.text.strip().replace("\n"," / ")[:60])
    print(f"slide {i}: layout={s.slide_layout.name!r}  text={txt[:4]}")
