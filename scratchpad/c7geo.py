from pptx import Presentation
p=Presentation("발표자료/파워포인트_템플릿.pptx")
for name in ["CUSTOM_7","CUSTOM_16","CUSTOM_15"]:
    lay=[l for l in p.slide_layouts if l.name==name][0]
    print(f"[{name}] placeholders:")
    for ph in lay.placeholders:
        try:
            print(f"   idx={ph.placeholder_format.idx} {str(ph.placeholder_format.type):16s} L={ph.left/914400:.2f} T={ph.top/914400:.2f} W={ph.width/914400:.2f} H={ph.height/914400:.2f}")
        except Exception as e:
            print(f"   idx={ph.placeholder_format.idx} (no geom)")
# also: template demo slide 3 (index2) — confirm layout + shapes look
print("\n=== 템플릿 데모 slide index2 (=3번 페이지) ===")
s=p.slides[2]
print("layout:", s.slide_layout.name)
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        print("  text:", repr(sh.text_frame.text.strip()[:50]))
