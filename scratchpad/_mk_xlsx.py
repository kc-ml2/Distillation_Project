import json, glob, openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from tensorboard.backend.event_processing.event_accumulator import EventAccumulator

METS = ["Bleu_1","Bleu_2","Bleu_3","Bleu_4","METEOR","ROUGE_L","CIDEr","SPICE"]

# rows: (label, group, zeroshot_dir_or_None, ttm_run_or_None)
ROWS = [
    ("small solo (증류 X)",          "target", "small_solo_ep19", None),
    ("small + LM 증류",              "target", "small_distill_ep19", None),
    ("small + ITC+LM 증류",          "target", None, "pt_smallreg_minilm_ttm_queue_lm"),
    ("base (우리, base_amp)",         "target", "base_amp_ep19", None),
    ("base_nodecay (우리)",          "ref",    "base_nodecay_ep19", None),
    ("BLIP-base 14M (공식)",         "ref",    "base", None),
    ("BLIP-large (티처, 공식)",       "ref",    "large", None),
]

def read_zeroshot(d):
    line = open(f"output/caption_zeroshot/{d}/evaluate.txt").read().strip().splitlines()[-1]
    j = json.loads(line)
    val  = {m: j.get(f"val_{m}")  for m in METS}
    test = {m: j.get(f"test_{m}") for m in METS}
    return val, test

def read_ttm(run):
    ev = sorted(glob.glob(f"output/{run}/tensorboard/**/events*", recursive=True))[-1]
    ea = EventAccumulator(ev, size_guidance={'scalars':0}); ea.Reload()
    val = {}
    for m in METS:
        tag = f"val_caption/{m}"
        val[m] = ea.Scalars(tag)[-1].value if tag in ea.Tags()['scalars'] else None
    return val, {m: None for m in METS}  # test not run in-loop

data = []
for label, grp, zd, ttm in ROWS:
    val, test = read_zeroshot(zd) if zd else read_ttm(ttm)
    data.append((label, grp, val, test))

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "caption ep19"
hdr_fill = PatternFill("solid", fgColor="2A78D6")
grp_fill = PatternFill("solid", fgColor="EEF4FD")
white = Font(color="FFFFFF", bold=True); bold = Font(bold=True)
thin = Side(style="thin", color="DEDCD6"); border = Border(thin,thin,thin,thin)
center = Alignment(horizontal="center")

# header
ws.cell(1,1,"모델 (ep19)"); ws.cell(1,2,"split")
for j,m in enumerate(METS): ws.cell(1,3+j,m)
for c in range(1,3+len(METS)):
    cell = ws.cell(1,c); cell.fill = hdr_fill; cell.font = white; cell.alignment = center; cell.border = border

r = 2
for label, grp, val, test in data:
    for split, vals in (("val", val), ("test", test)):
        if all(v is None for v in vals.values()):  # skip empty test row (ttm)
            continue
        ws.cell(r,1,label); ws.cell(r,2,split)
        if grp == "ref":
            ws.cell(r,1).fill = grp_fill; ws.cell(r,2).fill = grp_fill
        for j,m in enumerate(METS):
            v = vals[m]
            cell = ws.cell(r,3+j, round(v,4) if v is not None else "—")
            cell.alignment = center; cell.border = border
            if m == "CIDEr": cell.font = bold
        ws.cell(r,1).border = border; ws.cell(r,2).border = border
        r += 1

ws.column_dimensions['A'].width = 24; ws.column_dimensions['B'].width = 6
for col in "CDEFGHIJ": ws.column_dimensions[col].width = 9
ws.freeze_panes = "C2"

note = ws.cell(r+1,1,"주: ITC+LM은 in-loop val_caption(같은 COCO val+generate 프로토콜, test 미측정). base/large는 공식 체크포인트(우리 ep19 아님). SPICE는 ITC+LM만 기록.")
note.font = Font(italic=True, size=9, color="78766F")

out = "docs/experiments/caption_scores_ep19.xlsx"
wb.save(out); print("saved", out, "rows", r-2)
