from pptx import Presentation
p=Presentation("발표자료/발표.pptx")
for i,s in enumerate(p.slides):
    imgs=sum(1 for sh in s.shapes if sh.shape_type==13)
    tbls=sum(1 for sh in s.shapes if sh.has_table)
    title=""
    for sh in s.shapes:
        if sh.has_text_frame and sh==(s.shapes.title if s.shapes.title else None):
            title=sh.text_frame.text[:40]
    if not title:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                title=sh.text_frame.text.strip()[:40];break
    print(f"{i:2d} [{s.slide_layout.name:22s}] img={imgs} tbl={tbls} | {title!r}")
