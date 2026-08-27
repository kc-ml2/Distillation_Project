# -*- coding: utf-8 -*-
"""Build the presentation from the Google-Slides-exported template."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from PIL import Image
import os

ROOT = "/home/minwoo/Distillation_Project/발표자료"
TPL = os.path.join(ROOT, "파워포인트_템플릿.pptx")
OUT = os.path.join(ROOT, "발표.pptx")
A = os.path.join(ROOT, "assets")

prs = Presentation(TPL)
LAY = {l.name: l for l in prs.slide_layouts}

# ---- remove the demo slides: drop the relationship (orphaned part won't be
#      serialized since python-pptx only walks reachable parts), then the sldId ----
sldIdLst = prs.slides._sldIdLst
for sid in list(sldIdLst):
    prs.part.drop_rel(sid.get(qn('r:id')))
    sldIdLst.remove(sid)

EMU = 914400
def inch(v): return Emu(int(v*EMU))

def ph(slide, idx):
    for p in slide.placeholders:
        if p.placeholder_format.idx == idx:
            return p
    return None

def set_title(slide, text):
    t = ph(slide, 0)
    if t is not None:
        t.text_frame.text = text
    return t

def set_body(slide, idx, bullets, sizes=None):
    """bullets: list of (level, text)."""
    body = ph(slide, idx)
    tf = body.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, item in enumerate(bullets):
        lvl, txt = item if isinstance(item, tuple) else (0, item)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = txt
        p.level = lvl
        if sizes:
            for r in p.runs:
                r.font.size = Pt(sizes)
    return body

def add_slide(layout_name):
    return prs.slides.add_slide(LAY[layout_name])

def remove_ph(slide, idx):
    """Delete an unused placeholder so its prompt text doesn't show."""
    p = ph(slide, idx)
    if p is not None:
        p._element.getparent().remove(p._element)

def fit_box(img_path, box_l, box_t, box_w, box_h):
    """Return (l,t,w,h) in inches fitting the image aspect inside the box (centered)."""
    w, h = Image.open(img_path).size
    ar = w / h
    bar = box_w / box_h
    if ar > bar:  # image wider -> constrain width
        nw = box_w; nh = box_w / ar
    else:
        nh = box_h; nw = box_h * ar
    nl = box_l + (box_w - nw) / 2
    nt = box_t + (box_h - nh) / 2
    return nl, nt, nw, nh

def add_image(slide, img_path, box):
    l, t, w, h = fit_box(img_path, *box)
    return slide.shapes.add_picture(img_path, inch(l), inch(t), inch(w), inch(h))

def add_textbox(slide, l, t, w, h, lines, size=12, bold_first=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(inch(l), inch(t), inch(w), inch(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln; p.alignment = align
        for r in p.runs:
            r.font.size = Pt(size)
            if bold_first and i == 0: r.font.bold = True
    return tb

def add_table(slide, rows, l, t, w, h, header=True, font=10):
    nr = len(rows); nc = len(rows[0])
    gtab = slide.shapes.add_table(nr, nc, inch(l), inch(t), inch(w), inch(h)).table
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = gtab.cell(ri, ci)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(font)
                    if header and ri == 0: r.font.bold = True
    return gtab

# =========================================================
# 1. TITLE
s = add_slide("TITLE")
set_title(s, "BLIP 지식증류를 통한\nVision-Language 모델 경량화")
sub = ph(s, 1)
if sub is not None: sub.text_frame.text = "최민우 · KC-ML2"

# 2. 왜 경량화·왜 KD
s = add_slide("CUSTOM_7")
set_title(s, "왜 경량화, 왜 하필 KD인가")
set_body(s, 1, [
    "목표: KD를 통한 '범용' 모델 축소 — quantization보다 손이 가지만, 그만큼 실용적 가치를 세우는 연구",
    "Quantization은 이득·손실이 명확 / KD는 원하는 하드웨어 스펙에 맞춰 유연하게 압축 가능",
    "극한 최적화 = KD(앞단) + Quantization(뒷단) 결합으로 도달",
])

# 3. KD 원리 (복선)
s = add_slide("CUSTOM_7")
set_title(s, "지식증류(KD)의 원리")
set_body(s, 1, [
    "티처의 soft 출력(dark knowledge) + 정답(hard label)이 함께 학생을 가르친다",
    "가장 쉬운 건 '마지막 출력단' 증류 — 몸통(hidden)은 티처·학생 차원이 달라 증류가 어렵다",
    "▶ 그런데 마지막 단의 '값'만 넘기면 항상 잘 될까?  (뒤에서 뒤집힌다)",
])

# 4. 왜 BLIP
s = add_slide("CUSTOM_7")
set_title(s, "왜 BLIP인가 — 세 개의 출력 구조")
set_body(s, 1, [
    "VLM인데 역할이 셋: 이미지-텍스트 검색 / 매칭 / 캡셔닝",
    "역할마다 loss가 다르다 → ITC(contrastive) · ITM(binary) · LM(생성)",
    "→ 출력 구조가 서로 다른 증류를 한 모델 안에서 비교하는 '자연 실험장'",
])

# 5. BLIP 구조도 (CUSTOM_7 base + 수동 이미지/캡션)
s = add_slide("CUSTOM_7")
set_title(s, "BLIP 구조 한눈에")
remove_ph(s, 1)
add_image(s, os.path.join(A, "figures/blip_architecture.png"), (0.3, 1.0, 9.4, 3.6))
add_textbox(s, 0.3, 4.75, 9.4, 0.5,
            ["Vision Encoder + Text Encoder(self→cross→ff) · 3개 끝단: ITC(cls 유사도) / ITM(binary head) / LM(causal)"],
            size=11)

# 6. BLIP-Small 제원 (title + table + note)
s = add_slide("CUSTOM_7")
set_title(s, "우리가 만든 학생: BLIP-Small")
# shrink body to top note, table below
set_body(s, 1, [
    "티처 = BLIP-Large → 학생 = BLIP-Small (Base는 크기 사다리 맥락)",
    "DINOv3 + MiniLM: base와 구조가 가장 유사(MiniLM은 Bert-base와 층수 동일 → 코드 재사용)",
    "데이터: COCO(clean) + VG(noisy), 1 epoch ≈ 597k, ~5h",
], sizes=12)
add_table(s, [
    ["", "Vision", "Text"],
    ["Large (teacher)", "ViT-L 304M", "Bert-base 109M"],
    ["Base", "DeiT-base 86M", "Bert-base 109M"],
    ["Small (ours)", "DINOv3 21.6M", "MiniLM 33M"],
], l=0.3, t=3.55, w=6.2, h=1.6, font=11)

# 7. ITC 증류
s = add_slide("CUSTOM_7")
set_title(s, "ITC 증류 — Contrastive + Momentum Queue")
set_body(s, 1, [
    "MoCo 모멘텀 큐: online 이미지 feature @ 모멘텀 텍스트(+큐)로 로짓 구성",
    "타겟 = hard + 모멘텀(이미지@텍스트) + 티처(티처 큐)의 convex 결합",
    "정답은 정답에, 오답은 오답에 (in-batch + 큐 네거티브)",
    "발견: contrastive 증류는 '배치가 충분히 커야' 작동한다 (네거티브 수가 신호의 전제)",
])

# 8. ITM 증류 (하이라이트)
s = add_slide("CUSTOM_7")
set_title(s, "ITM 증류 — 값이 아니라 '구조'를 넘겨라 ★")
set_body(s, 1, [
    "ITM = 이미지-텍스트 매칭 확률 (cross-attention → binary head)",
    "문제: binary 로짓을 직접 증류하면 오히려 성능이 하락 (gradient가 죽거나 학습 상한이 생김)",
    "해결: 값이 아니라 1 pos + 4 neg의 '관계 구조'를 증류 → 확실한 이득  ◀ 앞의 복선 회수",
    "비용: pair 3B → 10B, +약 8GB (메모리·시간 때문에 관계 수 변화 실험은 못 함)",
])

# 9. LM 증류
s = add_slide("CUSTOM_7")
set_title(s, "LM 증류 — 단순 값 증류로 충분")
set_body(s, 1, [
    "causal LM, 티처의 next-token 분포와 학생의 KLdiv(T‖S)",
    "ITM과 정반대: 여기선 '값'만 넘겨도 잘 된다",
    "결과: 캡션 CIDEr가 전 에폭에서 distill > solo (ep19 기준 +2.4%)",
])

# 10. 통찰 (CUSTOM_7 base: 좌 text, 우 RKD 이미지)
s = add_slide("CUSTOM_7")
set_title(s, "핵심 통찰 — 출력 차원이 증류 방식을 가른다")
remove_ph(s, 1)
add_textbox(s, 0.3, 1.0, 4.5, 4.0, [
    "· LM(출력차원 큼, vocab): 값 증류 OK — dark knowledge 풍부, 후반까지 유효",
    "· ITM(출력차원 작음, binary): 값 증류 위험 → 구조(relation) 증류 필수",
    "· ITC(contrastive): relation이 본질 + 배치 크기가 전제",
    "→ 출력이 작을수록 '값' 대신 '구조'를 넘겨라",
], size=13)
add_image(s, os.path.join(A, "figures/rkd_figure1.png"), (4.95, 1.0, 4.55, 4.0))

# 11. 결과 (CUSTOM_7 base: 좌 곡선 이미지, 우 표+펀치라인)
s = add_slide("CUSTOM_7")
set_title(s, "결과 — 경량 학생이 base를 따라잡다")
remove_ph(s, 1)
add_image(s, os.path.join(A, "curves/itc_r_mean_overlay_zoom.png"), (0.25, 1.0, 4.55, 4.0))
add_table(s, [
    ["실험", "ITC r_mean", "ITM", "CIDEr"],
    ["base (목표)", "65.67", "73.22", "1.11"],
    ["small baseline", "63.10", "70.93", "1.08"],
    ["ITC", "65.22", "71.11", "—"],
    ["LM", "63.97", "71.60", "1.10"],
    ["ITC+LM", "66.11", "72.09", "1.10"],
    ["ITM", "64.19", "71.60", "—"],
], l=4.9, t=1.0, w=4.85, h=2.9, font=10)
add_textbox(s, 4.9, 4.05, 4.85, 1.0,
            ["ITC+LM 66.11 > base 65.67 > small 63.10", "→ 경량 학생이 retrieval에서 base 재현치를 넘음"],
            size=11, bold_first=True)

# 12. 의의/결론
s = add_slide("CUSTOM_7")
set_title(s, "의의 / 결론")
set_body(s, 1, [
    "KD는 '구조 전달'이 핵심 — 단순 값 증류가 통하는 loss(LM)도 있지만, binary 로짓 증류는 치명적",
    "Contrastive 증류는 배치가 충분히 커야 작동한다",
    "가설: 출력차원↑ → 값 증류(dark knowledge 풍부) / 출력차원↓ → relation 증류",
    "Relational KD는 티처·학생이 같은 목표를 가리켜 hidden 증류보다 안전",
    "향후: quantization 결합으로 극한 압축",
])

# 13. 감사/QA
s = add_slide("SECTION_HEADER")
set_title(s, "감사합니다 — Q&A")

# B1. collapse (백업) — CUSTOM_7 base
s = add_slide("CUSTOM_7")
set_title(s, "[백업] ITC+LM 후반 붕괴와 해결")
remove_ph(s, 1)
add_image(s, os.path.join(A, "curves/collapse_vs_fix.png"), (0.25, 1.0, 4.6, 3.9))
add_textbox(s, 5.0, 1.1, 4.5, 3.8, [
    "· ~505k step에서 r_mean 66→58 급락, ~600k까지 요동 후 회복",
    "· (15k tail 평균엔 안 잡힘 — 원본 곡선을 봐야 보인다)",
    "· 원인: 티처 온도가 너무 soft + 큐 중복행에 티처 과확신 grad",
    "· 해결: 티처 온도 sharpen(tctmp0.0157) → 붕괴 거의 소멸, 최종 66.11",
], size=12)

# B2. 캡션 정성 비교 (백업) — CUSTOM_7 base, 3행 이미지+4모델 캡션
s = add_slide("CUSTOM_7")
remove_ph(s, 1)
set_title(s, "[백업] 캡션 정성 비교 — 4개 모델 (ep19)")
cap_rows = [
    ("ex1_508917.jpg", "기차역 (distill이 baseline 이김)", [
        "baseline: a train traveling down train tracks next to a platform",
        "LM distill: a train on the tracks at a train station",
        "base: a train pulling into a train station next to a platform",
        "Large(teacher): a train on the tracks",
    ]),
    ("ex2_570834.jpg", "자전거 실은 객차 (distill이 장면 오인 정정)", [
        "baseline: a man standing in front of a bike rack",
        "LM distill: a man standing next to a bike in a train station",
        "base: a group of people standing next to a row of bikes",
        "Large(teacher): a man standing next to a bunch of bikes",
    ]),
    ("ex3_480720.jpg", "거울 속 고양이 (capacity 필요 — 작은 모델 둘 다 실패)", [
        "baseline: a couple of cats sitting on top of a window sill",
        "LM distill: a couple of cats sitting on a window sill",
        "base: a siamese cat looking at itself in a mirror",
        "Large(teacher): a cat looking at its reflection in a mirror",
    ]),
]
for i, (img, scene, caps) in enumerate(cap_rows):
    t = 1.0 + i * 1.48
    add_image(s, os.path.join(A, "captions", img), (0.3, t, 2.0, 1.35))
    add_textbox(s, 2.45, t - 0.05, 7.35, 1.45, [scene] + caps, size=9, bold_first=True)

prs.save(OUT)
print("saved", OUT, "| slides:", len(prs.slides._sldIdLst))
